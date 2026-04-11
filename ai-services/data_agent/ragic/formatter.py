"""
@file        formatter.py
@description Multi-format output formatter for Ragic query results.
             Supports JSON (default), CSV, and Excel (.xlsx).
@lastUpdate  2026-04-11 13:28:14
@author      Daniel Chung
@version     1.0.0
"""

import csv
import io
import logging

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from data_agent.ragic.models import FormattedRecord

logger = logging.getLogger(__name__)


def _build_table_rows(
    records: list[FormattedRecord],
    field_labels: dict[str, str],
) -> tuple[list[str], list[list[str]]]:
    """Extract ordered column headers and row data from formatted records.

    Returns:
        (headers, rows) where headers are human-readable labels (falling
        back to field_id) and rows contain string-coerced cell values.
    """
    seen_ids: dict[str, None] = {}
    for rec in records:
        for fid in rec.fields:
            if fid.startswith("_"):
                continue
            if fid not in seen_ids:
                seen_ids[fid] = None

    ordered_ids = list(seen_ids.keys())

    headers = ["ragic_id"]
    for fid in ordered_ids:
        headers.append(field_labels.get(fid, fid))

    rows: list[list[str]] = []
    for rec in records:
        row = [rec.ragic_id]
        for fid in ordered_ids:
            raw = rec.fields.get(fid, "")
            if isinstance(raw, list):
                row.append(", ".join(str(v) for v in raw))
            elif isinstance(raw, dict):
                row.append(str(raw))
            else:
                row.append(str(raw) if raw is not None else "")
        rows.append(row)

    return headers, rows


def to_csv_bytes(
    records: list[FormattedRecord],
    field_labels: dict[str, str],
) -> bytes:
    """Convert records to UTF-8 encoded CSV bytes with BOM for Excel compat."""
    headers, rows = _build_table_rows(records, field_labels)

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)

    csv_text = buf.getvalue()
    return b"\xef\xbb\xbf" + csv_text.encode("utf-8")


def to_excel_bytes(
    records: list[FormattedRecord],
    field_labels: dict[str, str],
    sheet_name: str = "Ragic Data",
) -> bytes:
    """Convert records to .xlsx bytes using openpyxl."""
    headers, rows = _build_table_rows(records, field_labels)

    wb = Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet()
    ws.title = sheet_name

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = cell.font.copy(bold=True)

    for row_idx, row_data in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row_data in rows[:50]:
            if col_idx - 1 < len(row_data):
                cell_len = len(str(row_data[col_idx - 1]))
                if cell_len > max_len:
                    max_len = cell_len
        adjusted_width = min(max_len + 2, 50)
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = adjusted_width

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
