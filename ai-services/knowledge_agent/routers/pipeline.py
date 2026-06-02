import os
from pathlib import Path
from typing import Any, Optional, cast

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import FileResponse
from knowledge_agent import service as kms_module
from pydantic import BaseModel
from shared.security import verify_internal_token

router = APIRouter(
    prefix="/pipeline",
    tags=["Knowledge Pipeline"],
    dependencies=[Depends(verify_internal_token)],
)


def _kms_check(root_id: str | None, user_role: str | None, operation: str) -> None:
    """Wrapper that translates KMS ValueError → HTTPException(403)."""
    try:
        kms = kms_module.get_knowledge_management_service()
        kms.check_access(root_id, user_role, operation)
    except ValueError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc

ARANGO_URL = os.getenv("ARANGO_URL", "http://127.0.0.1:8529")
ARANGO_DB = os.getenv("ARANGO_DB", "abc_desktop")
ARANGO_USER = os.getenv("ARANGO_USER", "root")
ARANGO_PASSWORD = os.getenv("ARANGO_PASSWORD", "")


class TriggerRequest(BaseModel):
    task: str
    file_id: str
    local_path: str
    root_id: str
    session_key: Optional[str] = None
    user_role: Optional[str] = None


@router.post("/vector")
async def trigger_vector(
    request: Request,
    file_id: str,
    root_id: str,
    user_role: Optional[str] = None,
) -> dict[str, object]:
    from celery_app.tasks import vectorize_task
    from kb_pipeline.arango_ops import ArangoOps

    user_role = request.headers.get("X-User-Role") or user_role
    _kms_check(root_id, user_role, operation="index")

    arango = ArangoOps()
    file_doc = arango.get_file(file_id)
    local_path = file_doc.get("local_path") if file_doc else None
    if not local_path:
        return {"error": "file not found or local_path missing"}
    result = cast(Any, vectorize_task).delay(file_id, local_path, root_id)
    arango.set_task_id(file_id, vector_task_id=result.id)
    return {
        "status": "queued",
        "file_id": file_id,
        "type": "vectorize",
        "task_id": result.id,
    }


@router.post("/trigger")
async def trigger_pipeline(request: Request, body: TriggerRequest) -> dict[str, object]:
    from celery_app.tasks import graph_task, vectorize_task
    from kb_pipeline.arango_ops import ArangoOps

    effective_role = request.headers.get("X-User-Role") or body.user_role
    _kms_check(body.root_id, effective_role, operation="index")

    arango = ArangoOps()
    vector_result = cast(Any, vectorize_task).delay(
        body.file_id, body.local_path, body.root_id, session_key=body.session_key
    )
    graph_result = cast(Any, graph_task).delay(
        body.file_id, body.local_path, session_key=body.session_key
    )
    arango.set_task_id(
        body.file_id, vector_task_id=vector_result.id, graph_task_id=graph_result.id
    )
    return {
        "status": "queued",
        "file_id": body.file_id,
        "vector_task_id": vector_result.id,
        "graph_task_id": graph_result.id,
    }


@router.post("/graph")
async def trigger_graph(
    request: Request,
    file_id: str,
    user_role: Optional[str] = None,
) -> dict[str, object]:
    from celery_app.tasks import graph_task
    from kb_pipeline.arango_ops import ArangoOps

    arango = ArangoOps()
    file_doc = arango.get_file(file_id)
    if not file_doc:
        raise HTTPException(status_code=404, detail="file not found")
    root_id = str(file_doc.get("knowledge_root_id", ""))
    if not root_id:
        raise HTTPException(status_code=400, detail="knowledge_root_id not found in file doc")

    role = request.headers.get("X-User-Role") or user_role
    _kms_check(root_id, role, operation="index")

    local_path = file_doc.get("local_path")
    result = cast(Any, graph_task).delay(file_id, local_path)
    arango.set_task_id(file_id, graph_task_id=result.id)
    return {
        "status": "queued",
        "file_id": file_id,
        "type": "graph",
        "task_id": result.id,
    }


@router.get("/vectors")
async def get_vectors(
    request: Request,
    file_id: str,
    limit: int = 50,
    offset: int = 0,
    user_role: Optional[str] = None,
) -> dict[str, object]:
    from kb_pipeline.arango_ops import ArangoOps
    from kb_pipeline.qdrant_ops import QdrantStore

    arango = ArangoOps()
    file_doc = arango.get_file(file_id)
    if not file_doc:
        raise HTTPException(status_code=404, detail="file not found")
    raw_root_id = file_doc.get("knowledge_root_id")
    root_id: str | None = str(raw_root_id) if raw_root_id else None
    if not root_id:
        return {"chunks": [], "total": 0, "file_id": file_id}
    role = request.headers.get("X-User-Role") or user_role
    _kms_check(root_id, role, operation="query")
    qdrant = QdrantStore()
    collection = f"knowledge_{root_id}"
    chunks = qdrant.get_chunks(collection, file_id, limit, offset)
    return {"chunks": chunks, "total": len(chunks), "file_id": file_id}


@router.get("/similar")
async def get_similar(
    request: Request,
    file_id: str,
    chunk_id: str,
    top_k: int = 10,
    user_role: Optional[str] = None,
) -> dict[str, object]:
    from kb_pipeline.arango_ops import ArangoOps
    from kb_pipeline.qdrant_ops import QdrantStore

    arango = ArangoOps()
    file_doc = arango.get_file(file_id)
    if not file_doc:
        raise HTTPException(status_code=404, detail="file not found")
    raw_root_id = file_doc.get("knowledge_root_id")
    root_id: str | None = str(raw_root_id) if raw_root_id else None
    if not root_id:
        return {"similar": []}
    role = request.headers.get("X-User-Role") or user_role
    _kms_check(root_id, role, operation="query")

    qdrant = QdrantStore()
    collection = f"knowledge_{root_id}"
    try:
        positive_id = int(chunk_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="chunk_id must be numeric")
    results = qdrant.recommend(collection, positive_id, limit=top_k)
    similar = []
    for r in results:
        pld = cast(dict[str, object], r.get("payload") or {})
        similar.append(
            {
                "chunk_id": str(r.get("id", "")),
                "text": str(pld.get("text_full") or pld.get("text") or ""),
                "score": cast(float, r.get("score", 0.0)),
            }
        )
    return {"similar": similar}


@router.post("/regenerate/{file_id}")
async def regenerate_pipeline(
    request: Request,
    file_id: str,
    user_role: Optional[str] = None,
) -> dict[str, object]:
    from celery_app.tasks import graph_task, vectorize_task
    from kb_pipeline.arango_ops import ArangoOps

    arango = ArangoOps()
    file_doc = arango.get_file(file_id)
    if not file_doc:
        raise HTTPException(status_code=404, detail="file not found")
    local_path = file_doc.get("local_path")
    root_id: str | None = str(file_doc.get("knowledge_root_id")) if file_doc.get("knowledge_root_id") else None
    if not local_path or not root_id:
        raise HTTPException(
            status_code=400,
            detail="file missing local_path or knowledge_root_id",
        )

    role = request.headers.get("X-User-Role") or user_role
    _kms_check(root_id, role, operation="index")

    vector_result = cast(Any, vectorize_task).delay(file_id, local_path, root_id)
    graph_result = cast(Any, graph_task).delay(file_id, local_path)
    arango.set_task_id(
        file_id, vector_task_id=vector_result.id, graph_task_id=graph_result.id
    )
    arango.update_status(file_id, vector_status="queued", graph_status="queued")
    return {
        "status": "queued",
        "file_id": file_id,
        "vector_task_id": vector_result.id,
        "graph_task_id": graph_result.id,
    }


@router.get("/graph")
async def get_graph(
    request: Request,
    file_id: str,
    user_role: Optional[str] = None,
) -> dict[str, object]:
    from kb_pipeline.arango_ops import ArangoOps

    arango = ArangoOps()
    file_doc = arango.get_file(file_id)
    if not file_doc:
        raise HTTPException(status_code=404, detail="file not found")
    raw_root_id = file_doc.get("knowledge_root_id")
    root_id: str | None = str(raw_root_id) if raw_root_id else None
    if root_id:
        role = request.headers.get("X-User-Role") or user_role
        _kms_check(root_id, role, operation="query")
    graph_data = arango.get_graph(file_id)
    return {"nodes": graph_data["nodes"], "edges": graph_data["edges"]}


@router.post("/retry")
async def retry_pipeline(
    request: Request,
    file_id: str,
    user_role: Optional[str] = None,
) -> dict[str, object]:
    from celery_app.tasks import graph_task, vectorize_task
    from kb_pipeline.arango_ops import ArangoOps

    arango = ArangoOps()
    file_doc = arango.get_file(file_id)
    if not file_doc:
        raise HTTPException(status_code=404, detail="file not found")

    local_path = file_doc.get("local_path", "")
    raw_root_id = file_doc.get("knowledge_root_id")
    root_id: str | None = str(raw_root_id) if raw_root_id else None
    if not local_path:
        raise HTTPException(status_code=400, detail="local_path missing")

    if root_id:
        role = request.headers.get("X-User-Role") or user_role
        _kms_check(root_id, role, operation="index")

    v_result = cast(Any, vectorize_task).delay(file_id, local_path, root_id)
    g_result = cast(Any, graph_task).delay(file_id, local_path)
    arango.set_task_id(file_id, vector_task_id=v_result.id, graph_task_id=g_result.id)
    arango.update_status(file_id, vector_status="pending", graph_status="pending")
    return {
        "status": "queued",
        "file_id": file_id,
        "vector_task_id": v_result.id,
        "graph_task_id": g_result.id,
    }


@router.get("/active-tasks")
async def get_active_celery_tasks() -> dict[str, object]:
    from celery_app.app import app as celery_app

    try:
        inspector = celery_app.control.inspect()
        active = inspector.active() or {}
        reserved = inspector.reserved() or {}

        task_ids: list[str] = []
        for worker_tasks in list(active.values()) + list(reserved.values()):
            for task in worker_tasks:
                if task_id := task.get("id"):
                    task_ids.append(task_id)

        return {"active_task_ids": task_ids}
    except Exception:
        return {"active_task_ids": []}


@router.post("/abort")
async def abort_pipeline(
    request: Request,
    file_id: str,
    user_role: Optional[str] = None,
) -> dict[str, object]:
    from celery_app.app import app as celery_app
    from kb_pipeline.arango_ops import ArangoOps

    arango = ArangoOps()
    file_doc = arango.get_file(file_id)
    if not file_doc:
        raise HTTPException(status_code=404, detail="file not found")

    raw_root_id = file_doc.get("knowledge_root_id")
    root_id: str | None = str(raw_root_id) if raw_root_id else None
    if root_id:
        role = request.headers.get("X-User-Role") or user_role
        _kms_check(root_id, role, operation="delete")

    revoked: list[str] = []
    vector_task_id = file_doc.get("vector_task_id")
    graph_task_id = file_doc.get("graph_task_id")

    for tid in [vector_task_id, graph_task_id]:
        if tid and isinstance(tid, str):
            celery_app.control.revoke(tid, terminate=True)
            revoked.append(tid)

    arango.update_status(
        file_id,
        vector_status="aborted",
        graph_status="aborted",
        failed_reason="任務已被使用者中止",
    )
    return {"status": "aborted", "file_id": file_id, "revoked": revoked}


@router.post("/delete")
async def delete_file_data(
    request: Request,
    file_id: str,
    user_role: Optional[str] = None,
) -> dict[str, object]:
    import base64
    import json

    import redis as redis_lib
    from celery_app.app import REDIS_URL
    from celery_app.app import app as celery_app
    from kb_pipeline.arango_ops import ArangoOps
    from kb_pipeline.qdrant_ops import QdrantStore

    arango = ArangoOps()
    file_doc = arango.get_file(file_id)
    if not file_doc:
        raise HTTPException(status_code=404, detail="file not found")

    raw_root_id = file_doc.get("knowledge_root_id")
    root_id: str | None = str(raw_root_id) if raw_root_id else None
    if root_id:
        role = request.headers.get("X-User-Role") or user_role
        _kms_check(root_id, role, operation="delete")

    revoked: list[str] = []
    try:
        for tid_key in ("vector_task_id", "graph_task_id"):
            tid = file_doc.get(tid_key)
            if tid and isinstance(tid, str):
                celery_app.control.revoke(tid, terminate=True)
                revoked.append(tid)

        r = redis_lib.from_url(REDIS_URL)
        queue_key = "celery"
        queue_len = cast(int, r.llen(queue_key))
        if queue_len and queue_len > 0:
            to_remove: list[str] = []
            queued_items = cast(list[bytes | str], r.lrange(queue_key, 0, queue_len - 1) or [])
            for raw in queued_items:
                try:
                    msg = json.loads(raw)
                    body = msg.get("body")
                    if isinstance(body, str):
                        body = json.loads(base64.b64decode(body))
                    args = body if isinstance(body, list) else (body or {}).get("args", [])
                    if isinstance(args, (list, tuple)) and len(args) > 0 and args[0] == file_id:
                        task_id = msg.get("headers", {}).get("id", "")
                        if task_id:
                            celery_app.control.revoke(task_id, terminate=True)
                            revoked.append(task_id)
                        to_remove.append(raw.decode() if isinstance(raw, bytes) else raw)
                except Exception:
                    continue
            for item in to_remove:
                r.lrem(queue_key, 1, item)
    except Exception:
        pass

    qdrant_deleted = False
    if root_id:
        try:
            qdrant = QdrantStore()
            qdrant.delete_by_file(f"knowledge_{root_id}", file_id)
            qdrant_deleted = True
        except Exception:
            pass

    arango_removed = arango.delete_file_data(file_id)

    local_deleted = False
    local_path = str(file_doc.get("local_path", ""))
    if local_path and Path(local_path).exists():
        try:
            Path(local_path).unlink()
            local_deleted = True
        except Exception:
            pass

    import os

    seaweed_deleted = False
    s3_path = str(file_doc.get("s3_path", ""))
    if s3_path:
        try:
            import httpx

            seaweed_base = os.getenv("SEAWEED_AIBOX_URL", "http://localhost:8888")
            seaweed_user = os.getenv("SEAWEED_USER", "")
            seaweed_pass = os.getenv("SEAWEED_PASS", "")
            async with httpx.AsyncClient(timeout=15.0) as client:
                resp = await client.request(
                    "DELETE",
                    f"{seaweed_base}/{s3_path}",
                    auth=(seaweed_user, seaweed_pass),
                )
                seaweed_deleted = resp.status_code < 400
        except Exception:
            pass

    return {
        "status": "deleted",
        "file_id": file_id,
        "revoked_tasks": revoked,
        "qdrant_deleted": qdrant_deleted,
        "arango_removed": arango_removed,
        "local_deleted": local_deleted,
        "seaweed_deleted": seaweed_deleted,
    }


@router.get("/logs")
async def get_pipeline_logs(
    request: Request,
    file_id: str,
    user_role: Optional[str] = None,
) -> dict[str, object]:
    from kb_pipeline.arango_ops import ArangoOps

    arango = ArangoOps()
    file_doc = arango.get_file(file_id)
    if file_doc:
        raw_root_id = file_doc.get("knowledge_root_id")
        root_id: str | None = str(raw_root_id) if raw_root_id else None
        if root_id:
            role = request.headers.get("X-User-Role") or user_role
            _kms_check(root_id, role, operation="query")
    logs = arango.get_job_logs(file_id)
    return {"file_id": file_id, "logs": logs, "count": len(logs)}


@router.get("/preview")
async def get_preview(
    request: Request,
    file_id: str,
    user_role: Optional[str] = None,
) -> dict[str, object]:
    from kb_pipeline.arango_ops import ArangoOps

    arango = ArangoOps()
    file_doc = arango.get_file(file_id)
    if not file_doc:
        raise HTTPException(status_code=404, detail="file not found")

    raw_root_id = file_doc.get("knowledge_root_id")
    root_id: str | None = str(raw_root_id) if raw_root_id else None
    if root_id:
        role = request.headers.get("X-User-Role") or user_role
        _kms_check(root_id, role, operation="query")

    local_path = str(file_doc.get("local_path"))
    if not local_path or not Path(local_path).exists():
        raise HTTPException(status_code=404, detail="file not found on disk")

    ext = Path(local_path).suffix.lower()
    content_type_map = {
        ".md": "markdown",
        ".txt": "text",
        ".pdf": "text",
        ".csv": "table",
        ".xlsx": "table",
        ".xls": "table",
    }
    preview_type = content_type_map.get(ext, "binary")

    if preview_type == "markdown":
        text = Path(local_path).read_text(encoding="utf-8", errors="replace")
        return {"file_id": file_id, "type": "markdown", "content": text}

    if ext == ".pdf":
        download_url = f"/pipeline/download?file_id={file_id}"
        return {"file_id": file_id, "type": "pdf_url", "url": download_url}

    if preview_type == "text":
        text = Path(local_path).read_text(encoding="utf-8", errors="replace")
        return {"file_id": file_id, "type": "text", "content": text[:5000]}

    if preview_type == "table":
        rows: list[dict[str, str | int | float]] = []
        headers: list[str] = []
        if ext in (".xlsx", ".xls"):
            import openpyxl

            wb = openpyxl.load_workbook(local_path, data_only=True)
            ws = wb.active
            if ws is None:
                raise HTTPException(status_code=500, detail="worksheet not found")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    headers = [str(c) if c is not None else "" for c in row]
                else:
                    rows.append(
                        {
                            str(headers[j]) if j < len(headers) else f"col{j}": str(c)
                            if c is not None
                            else ""
                            for j, c in enumerate(row)
                        }
                    )
            return {
                "file_id": file_id,
                "type": "table",
                "headers": headers,
                "rows": rows[:200],
            }
        if ext == ".csv":
            import csv

            with open(local_path, encoding="utf-8", errors="replace") as f:
                reader = csv.DictReader(f)
                for i, row in enumerate(reader):
                    if i == 0:
                        headers = list(row.keys())
                    rows.append(row)
                    if i >= 199:
                        break
            return {
                "file_id": file_id,
                "type": "table",
                "headers": headers,
                "rows": rows,
            }

    if ext in (".docx", ".doc"):
        from docx import Document

        doc = Document(local_path)
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables[:50]:
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                line = " | ".join(cells).strip()
                if line:
                    rows.append(line)
            if rows:
                parts.append("\n".join(rows))
        return {
            "file_id": file_id,
            "type": "text",
            "content": "\n".join(parts[:200]),
        }

    return {"file_id": file_id, "type": "binary", "message": "不支援的檔案格式"}


@router.get("/download")
async def download_file(
    request: Request,
    file_id: str,
    user_role: Optional[str] = None,
) -> FileResponse:
    from kb_pipeline.arango_ops import ArangoOps

    arango = ArangoOps()
    file_doc = arango.get_file(file_id)
    if not file_doc:
        raise HTTPException(status_code=404, detail="file not found")

    raw_root_id = file_doc.get("knowledge_root_id")
    root_id: str | None = str(raw_root_id) if raw_root_id else None
    if root_id:
        role = request.headers.get("X-User-Role") or user_role
        _kms_check(root_id, role, operation="query")

    local_path: str = str(file_doc.get("local_path"))
    if not local_path or not Path(local_path).exists():
        raise HTTPException(status_code=404, detail="file not found on disk")

    filename: str = str(file_doc.get("filename", "download"))
    file_type: str = str(file_doc.get("file_type", "application/octet-stream"))
    import mimetypes

    guessed = mimetypes.guess_type(filename)
    mime: str = guessed[0] if guessed[0] else str(file_type)

    return FileResponse(
        path=local_path,
        filename=filename,
        media_type=mime,
        content_disposition_type="inline",
    )
